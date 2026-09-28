from io import BytesIO

from docx import Document


def generar_acta(concurso: dict, postor: dict, resultado_evaluacion: dict) -> BytesIO:
    dictamen = resultado_evaluacion["dictamen"]

    doc = Document()
    doc.add_heading("Acta de Dictamen Técnico", level=1)
    doc.add_paragraph(f"Concurso: {concurso['nombre']}")
    doc.add_paragraph(f"Postor: {postor['razon_social']} (RUC {postor['ruc']})")
    doc.add_paragraph(f"Veredicto final: {dictamen['veredicto']}")
    doc.add_paragraph(f"Puntaje técnico: {dictamen['puntaje']}")
    doc.add_paragraph(dictamen["justificacion"])

    doc.add_heading("Evaluación Legal", level=2)
    for r in resultado_evaluacion["resultados_legal"]:
        doc.add_paragraph(f"{r['criterio_descripcion']} — {r['veredicto']}")
        doc.add_paragraph(r["justificacion"])

    doc.add_heading("Evaluación Técnica", level=2)
    for r in resultado_evaluacion["resultados_tecnico"]:
        doc.add_paragraph(f"{r['criterio_descripcion']} — {r['veredicto']}")
        doc.add_paragraph(r["justificacion"])

    doc.add_heading("Citas y Trazabilidad", level=2)
    for cita in dictamen["citas"]:
        doc.add_paragraph(f"Página {cita['pagina']}: {cita['extracto']}")

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer
